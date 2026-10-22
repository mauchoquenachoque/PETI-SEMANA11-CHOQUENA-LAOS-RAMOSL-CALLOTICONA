# -*- coding: utf-8 -*-
"""
Script: construir_word_final.py
Propósito: Construir el documento Word oficial del Taller Semana 11
           basado en SI886-PLANTILLA-TALLER.docx, con carátula UPT,
           todas las secciones requeridas, tablas estilizadas, capturas
           de Metabase y gráficos de Python, para el Caso 07:
           INSTITUTO SUPERIOR TECNOLÓGICO VERTIENTE SUR.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PLANTILLA_PATH = r"C:\Users\HP\Downloads\SI886-PLANTILLA-TALLER.docx"
OUTPUT_DOCX = os.path.join(SCRIPT_DIR, "SI886-S11-TALLER-Grupo07.docx")

def set_cell_background(cell, hex_color):
    """Establece el color de fondo de una celda de tabla."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos de una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table(table, header_bg="1F4E79", alt_bg="F8FAFC", border_color="D1D5DB"):
    """Aplica formato profesional a una tabla en Word."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Encabezado
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8.5)
                run.font.name = "Calibri"
    # Filas de datos
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 0 else "FFFFFF"
        for cell in row.cells:
            if bg != "FFFFFF":
                set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    run.font.name = "Calibri"

def add_styled_heading(doc, text, level):
    """Agrega un encabezado estilizado con la paleta UPT."""
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0] if h.runs else h.add_run(text)
    run.font.name = "Calibri"
    if level == 1:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(31, 78, 121) # #1F4E79
    elif level == 2:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 44, 89) # #0F2C59
    elif level == 3:
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(55, 65, 81) # #374151
    return h

def add_styled_paragraph(doc, text, bold_prefix=None, space_after=4):
    """Agrega un párrafo justificado con estilo profesional."""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    """Agrega una viñeta estilizada con sangría sin requerir estilo List Bullet."""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = "Calibri"
    r_bullet.font.size = Pt(9.5)
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = RGBColor(31, 78, 121)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_figure_with_caption(doc, img_path, caption_text, num_fig, width_in=6.2):
    """Inserta una figura centrada con número y leyenda formal."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap_lbl = p_cap.add_run(f"Figura {num_fig}. ")
        r_cap_lbl.font.name = "Calibri"
        r_cap_lbl.font.size = Pt(8.5)
        r_cap_lbl.font.bold = True
        r_cap_lbl.font.color.rgb = RGBColor(31, 78, 121)

        r_cap_txt = p_cap.add_run(caption_text)
        r_cap_txt.font.name = "Calibri"
        r_cap_txt.font.size = Pt(8.5)
        r_cap_txt.font.italic = True
        r_cap_txt.font.color.rgb = RGBColor(75, 85, 99)
    else:
        print(f"[WARN] No se encontró la imagen: {img_path}")

print("Iniciando construcción del informe Word final basado en la plantilla UPT...")

doc = docx.Document(PLANTILLA_PATH)

# 1. Configurar Carátula
# P10: TALLER
# P11: Título del taller
p11 = doc.paragraphs[11]
p11.text = "DIAGNÓSTICO DE GOBIERNO DIGITAL Y TABLERO DE LÍNEA BASE\nEMPRESA 07: INSTITUTO SUPERIOR TECNOLÓGICO VERTIENTE SUR"
for r in p11.runs:
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)

# P14: Curso
p14 = doc.paragraphs[14]
p14.text = "SI-886 · PLANEAMIENTO ESTRATÉGICO DE TI"
for r in p14.runs:
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True

# P17: Autores
p17 = doc.paragraphs[17]
p17.text = ("Laos González, Adriana Rafaela — Código: 2023077474\n"
            "Ramos Loza, Mariela Estefany — Código: 2023077478\n"
            "Calloticona Chambilla, Marymar Danytza — Código: 2023076791\n"
            "Choqueña Choque, Mauricio Arian — Código: 2023076799")
for r in p17.runs:
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True

# P20: Docente
p20 = doc.paragraphs[20]
p20.text = "DR. OSCAR JUAN JIMENEZ FLORES"
for r in p20.runs:
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True

# P27: Tacna - Perú
p27 = doc.paragraphs[27]
p27.text = "TACNA — PERÚ\n2026"
for r in p27.runs:
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True

# 2. Limpiar página "Cómo se usa esta plantilla" y el esqueleto vacío posterior a la carátula
# Mantendremos la carátula intacta (párrafos 0 a 28) y reconstruiremos el cuerpo completo con máxima rigurosidad
# Eliminamos los elementos de cuerpo desde el índice hacia abajo para repoblar con el contenido completo y formateado
body = doc.element.body
total_children = len(body)
# Borramos los párrafos y tablas desde el elemento 29 hasta el final, preservando sectPr
for child in list(body)[29:]:
    tag = child.tag.split('}')[-1]
    if tag != 'sectPr':
        body.remove(child)

print("Carátula preservada y espacio de trabajo listo para repoblación...")

# =========================================================================
# ÍNDICE GENERAL
# =========================================================================
add_styled_heading(doc, "Índice General", level=1)
toc_items = [
    ("1. Información sobre el evento práctico", "3"),
    ("    1.1 Título del evento práctico", "3"),
    ("    1.2 Objetivos", "3"),
    ("    1.3 Tiempo de duración", "3"),
    ("    1.4 Resultados de aprendizaje", "3"),
    ("    1.5 Recursos utilizados", "4"),
    ("    1.6 Normas de seguridad y confidencialidad", "4"),
    ("2. Procedimiento o metodología", "5"),
    ("    2.1 Paso A — Catálogo de servicios y su nivel de digitalización", "5"),
    ("    2.2 Paso B — Diagnóstico de los ocho ejes con línea base", "8"),
    ("    2.3 Paso C — Despliegue del tablero de línea base en Metabase", "12"),
    ("    2.4 Paso D — Redacción de la Sección 6.1 del PETI", "16"),
    ("    2.5 Paso E — Validación cruzada y comprobaciones de integridad", "21"),
    ("    2.6 Paso F — Control de versiones y transferencia técnica", "22"),
    ("3. Resultados y evidencias procedimentales", "23"),
    ("    3.1 Los tres resultados que se califican", "23"),
    ("    3.2 Lista de comprobación del taller (13 verificaciones)", "24"),
    ("    3.3 Evaluación de la rúbrica procedimental (20 puntos)", "26"),
    ("    3.4 Problemas identificados y oportunidades de mejora", "27"),
    ("4. Conclusiones estratégicas y técnicas", "28"),
    ("5. Cuestionario de evaluación técnica", "29"),
    ("6. Referencias bibliográficas (Normas APA)", "32"),
    ("7. Anexos documentales, analíticos y gráficos", "33")
]
for item, pag in toc_items:
    p_t = doc.add_paragraph()
    p_t.paragraph_format.line_spacing = 1.05
    p_t.paragraph_format.space_after = Pt(2)
    r1 = p_t.add_run(item)
    r1.font.name = "Calibri"
    r1.font.size = Pt(9)
    if not item.startswith("    "):
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(31, 78, 121)
    else:
        r1.font.color.rgb = RGBColor(55, 65, 81)
    r_dots = p_t.add_run(" " + "." * (75 - len(item)) + " ")
    r_dots.font.color.rgb = RGBColor(156, 163, 175)
    r2 = p_t.add_run(pag)
    r2.font.bold = True
    r2.font.size = Pt(9)

doc.add_page_break()

# =========================================================================
# 1. INFORMACIÓN SOBRE EL EVENTO PRÁCTICO
# =========================================================================
add_styled_heading(doc, "1. Información sobre el evento práctico", level=1)

add_styled_heading(doc, "1.1 Título del evento práctico", level=2)
add_styled_paragraph(doc, "Diagnóstico de gobierno digital, formulación del catálogo de servicios, evaluación de línea base en ocho ejes y construcción del tablero de supervisión en Metabase para el Instituto Superior Tecnológico Vertiente Sur (Caso 07).")

add_styled_heading(doc, "1.2 Objetivos", level=2)
add_bullet_point(doc, "Levantar el catálogo formal de servicios de la institución, categorizando su nivel actual de digitalización bajo el estándar de 0 a 5 y proyectando las metas del PETI.", "• Objetivo 1: ")
add_bullet_point(doc, "Evaluar con rigurosidad empírica los ocho ejes del gobierno digital peruano, sustentando cada puntaje en evidencias documentales verificables.", "• Objetivo 2: ")
add_bullet_point(doc, "Establecer la línea base numérica de cada indicador institucional o declarar explícitamente el estado 'No se mide' para los procesos ciegos.", "• Objetivo 3: ")
add_bullet_point(doc, "Determinar el nivel de madurez global de gobierno digital del instituto y ubicarlo en el modelo estándar de cinco niveles.", "• Objetivo 4: ")
add_bullet_point(doc, "Aplicar la prueba del principio de 'una sola vez' (Once-Only Principle), cuantificando la redundancia documental que sufren los estudiantes y docentes.", "• Objetivo 5: ")
add_bullet_point(doc, "Desplegar el tablero analítico de línea base en Metabase sobre un contenedor PostgreSQL local, construyendo al menos seis visualizaciones interactivas.", "• Objetivo 6: ")
add_bullet_point(doc, "Redactar la Sección 6.1 (Situación actual del gobierno digital) del PETI, sintetizando las brechas de mayor impacto que alimentarán los objetivos de la Sección 6.2.", "• Objetivo 7: ")

add_styled_heading(doc, "1.3 Tiempo de duración", level=2)
add_styled_paragraph(doc, "100 minutos cronológicos en sesión síncrona de laboratorio de cómputo, complementados con horas de trabajo de campo, análisis de datos y redacción técnica en equipo.")

add_styled_heading(doc, "1.4 Resultados de aprendizaje", level=2)
add_styled_paragraph(doc, "Al concluir el presente evento práctico, el equipo de estudiantes de Ingeniería de Sistemas está plenamente capacitado para:", bold_prefix="Competencia Procedimental: ")
add_bullet_point(doc, "Auditar el ecosistema de servicios tecnológicos de una organización de educación superior privada, diferenciando el promedio simple del promedio ponderado por volumen transaccional.")
add_bullet_point(doc, "Aplicar la normativa técnica de la Secretaría de Gobierno y Transformación Digital (PCM/SEGDI) adaptándola a entidades educativas privadas reguladas por MINEDU, SUNEDU, ANPD y SUNAT.")
add_bullet_point(doc, "Formular proyectos de instrumentación técnica para indicadores ciegos, comprendiendo que sin línea base resulta inviable demostrar valor público o retorno de inversión en TI.")
add_bullet_point(doc, "Diseñar y desplegar cuadros de mando integrados (BI Dashboards) mediante contenedores Docker, bases de datos relacionales y herramientas de visualización ejecutiva.")

add_styled_heading(doc, "1.5 Recursos utilizados", level=2)
add_styled_paragraph(doc, "Para la ejecución del taller y la construcción de los artefactos de auditoría se emplearon los siguientes recursos tecnológicos:")

# Tabla 0 de recursos
recursos_data = [
    ["Recurso", "Versión", "Para qué se usó en el taller"],
    ["Docker Engine / Desktop", "v26.1 / Linux", "Despliegue de los contenedores locales aislados de PostgreSQL y Metabase."],
    ["PostgreSQL", "16.2-alpine", "Base de datos relacional analítica receptora de las 4 tablas del taller (puerto 5433)."],
    ["Metabase BI Platform", "v0.49.5", "Construcción del tablero de línea base con las 6 visualizaciones ejecutivas (puerto 3001)."],
    ["Python / pandas / matplotlib", "Python 3.14.0", "Cálculo matemático de brechas ponderadas, redundancia documental y radares polares."],
    ["SQLAlchemy & psycopg2", "v2.1.6 / 2.9.13", "Conexión programática para la ingesta y persistencia automatizada de CSVs hacia PostgreSQL."],
    ["Plantilla Oficial UPT", "SI886-PLANTILLA", "Estructuración del informe pericial formal bajo estándares de calidad académica EPIS."]
]
tbl_rec = doc.add_table(rows=len(recursos_data), cols=3)
for r_i, row in enumerate(recursos_data):
    for c_i, val in enumerate(row):
        tbl_rec.rows[r_i].cells[c_i].paragraphs[0].text = val
style_table(tbl_rec)

add_styled_heading(doc, "1.6 Normas de seguridad y confidencialidad", level=2)
add_bullet_point(doc, "Clasificación de la Información: El catálogo de servicios, volúmenes de transacción y debilidades de control del Instituto Superior Tecnológico Vertiente Sur son tratados bajo estricta reserva institucional y propósitos académicos.", "1. Confidencialidad: ")
add_bullet_point(doc, "Seguridad de Datos Personales: Todo el procesamiento analítico se ejecutó sobre datos agregados en el entorno local (127.0.0.1); no se extrajeron ni almacenaron datos identificatorios de los 3 240 alumnos ni del personal en planilla.", "2. Privacidad (Ley 29733): ")
add_bullet_point(doc, "Tratamiento de Vulnerabilidades: Los incidentes críticos (como la alteración no autorizada de 17 actas de notas) se documentan a nivel de proceso y control de acceso, sin publicar vectores de ataque explotables.", "3. Seguridad Digital: ")
add_bullet_point(doc, "Consentimiento de Observación: La auditoría del recorrido del usuario se efectuó observando las interfaces de usuario estándar de los sistemas en producción sin alterar registros académicos.", "4. Integridad: ")

doc.add_page_break()

# =========================================================================
# 2. PROCEDIMIENTO O METODOLOGÍA
# =========================================================================
add_styled_heading(doc, "2. Procedimiento o metodología", level=1)
add_styled_paragraph(doc, "La metodología aplicada sigue rigurosamente los seis pasos procedurales establecidos en la guía técnica del taller SI-886, articulando los datos del caso de estudio institucional con las herramientas de software desarrolladas.")

# -------------------------------------------------------------------------
# PASO A
# -------------------------------------------------------------------------
add_styled_heading(doc, "2.1 Paso A — Catálogo de servicios y su nivel de digitalización", level=2)
add_styled_paragraph(doc, "A partir de la Ficha Técnica Institucional (DOCS/FICHA.md), el extracto de la Mesa de Servicio con 2 299 tickets anuales (DOCS/mesa-de-servicio-y-sla.md) y los procesos misionales de educación superior, se identificaron y estructuraron diez (10) servicios clave. Cada servicio fue categorizado bajo la escala estándar de seis niveles de madurez digital (Nivel 0: Nulo/Analógico a Nivel 5: Proactivo/Predictivo).")

add_styled_paragraph(doc, "El archivo generado en el repositorio corresponde a 06_gobierno_digital/GD01_catalogo_servicios.csv, cuya transcripción íntegra se presenta a continuación:")

# Tabla del Catálogo de Servicios
cat_headers = ["ID", "Servicio", "Usuario", "Volumen", "Act.", "Doc.", "Red.", "Meta", "Justificación Estratégica"]
cat_rows = [
    cat_headers,
    ["S-01", "Matrícula académica semestral", "Estudiante", "6,480", "1", "4", "3", "4", "Servicio troncal de mayor fricción; 100% digital elimina colas."],
    ["S-02", "Registro y publicación de actas", "Docente/Est.", "45,360", "2", "2", "2", "4", "Firma digital; erradica vulnerabilidad de alteración ilícita."],
    ["S-03", "Pago pensiones y comprobantes", "Est./Apoderado", "32,400", "1", "2", "1", "4", "Pasarela online y conciliación contable en tiempo real."],
    ["S-04", "Aula virtual y entrega tareas", "Estudiante/Doc.", "129,600", "3", "0", "0", "4", "Plataforma crítica; sincronización LDAP/SSO y nube."],
    ["S-05", "Emisión constancias y notas", "Estudiante/Egr.", "2,100", "1", "3", "3", "4", "Autoservicio con QR; reduce plazo de 10 días a 2 minutos."],
    ["S-06", "Trámite titulación profesional", "Egresado", "480", "1", "7", "5", "3", "Expediente digital y seguimiento en línea de dictámenes."],
    ["S-07", "Admisión nuevos postulantes", "Postulante ext.", "1,450", "1", "4", "0", "4", "Inscripción online con validación RENIEC y pago integrado."],
    ["S-08", "Restablecimiento credenciales", "Comunidad", "2,299", "0", "1", "1", "4", "Autoservicio 2FA/MFA; descongestiona 30% de llamadas."],
    ["S-09", "Préstamo biblioteca y catálogo", "Estudiante/Doc.", "8,500", "1", "1", "1", "4", "Catálogo OPAC web y biblioteca digital para las 7 carreras."],
    ["S-10", "Licencias y boletas de pago", "Personal/Doc.", "2,640", "1", "2", "1", "4", "Portal empleado con firma electrónica; cumple norma laboral."]
]
tbl_cat = doc.add_table(rows=len(cat_rows), cols=9)
for r_i, row in enumerate(cat_rows):
    for c_i, val in enumerate(row):
        tbl_cat.rows[r_i].cells[c_i].paragraphs[0].text = val
style_table(tbl_cat)

add_styled_paragraph(doc, "Para procesar cuantitativamente el catálogo se programó y ejecutó el script 06_gobierno_digital/GD02_analisis_servicios.py, el cual generó el siguiente reporte pericial y la visualización gráfica:")

# Salida textual de GD02
with open(os.path.join(SCRIPT_DIR, "evidencias", "S11", "salidas", "GD02_salida.txt"), "r", encoding="utf-8") as f:
    gd02_txt = f.read()

p_log1 = doc.add_paragraph()
p_log1.paragraph_format.left_indent = Inches(0.2)
p_log1.paragraph_format.right_indent = Inches(0.2)
p_log1.paragraph_format.space_before = Pt(4)
p_log1.paragraph_format.space_after = Pt(6)
r_log1 = p_log1.add_run(gd02_txt)
r_log1.font.name = "Consolas"
r_log1.font.size = Pt(7.5)
r_log1.font.color.rgb = RGBColor(30, 41, 59)
set_cell_background(tbl_cat.rows[0].cells[0], "1F4E79") # dummy call to keep style

# Análisis de contraste simple vs ponderado
add_styled_paragraph(doc, "El análisis comparativo evidencia una severa discrepancia entre métricas:", bold_prefix="Contraste de Nivel Medio Simple vs. Ponderado: ")
add_bullet_point(doc, "El promedio aritmético simple de los 10 servicios arroja un valor de 1.20 de 5.00, reflejando que la abrumadora mayoría de servicios de ventanilla son netamente analógicos o informativos básicos.")
add_bullet_point(doc, "El promedio ponderado por volumen transaccional anual (231 309 transacciones) alcanza 2.31 de 5.00. Este aparente incremento numérico se explica por la altísima frecuencia de interacción en el Aula Virtual Moodle (129 600 accesos en Nivel 3).")
add_bullet_point(doc, "Veredicto del Diagnóstico: El promedio ponderado desvela un 'falso confort'. Si se aíslan las consultas pasivas en el Moodle, los servicios transaccionales que definen la retención, recaudación y legalidad institucional (matrícula, cobranza, emisión de constancias y mesa de ayuda) están paralizados en Nivel 1 y Nivel 0, exigiendo colas físicas y trámites en papel.")

# Prueba de una sola vez
add_styled_paragraph(doc, "En el catálogo se auditaron 26 requisitos documentales exigidos a los usuarios. De ellos, 17 documentos ya constan en los archivos físicos o bases de datos del instituto (DNI registrado, fichas de matrícula, constancias de no adeudo generadas por tesorería, certificados modulares previos). Esto representa un Índice de Redundancia Documental del 65.38%, demostrando que 2 de cada 3 documentos solicitados representan una carga burocrática injustificada impuesta al estudiante.", bold_prefix="Prueba del Principio de 'Una Sola Vez': ")

# Ranking de Brechas
add_styled_paragraph(doc, "Al ponderar la brecha unitaria (Nivel Objetivo - Nivel Actual) por el volumen anual de atenciones, se identificó que el Top 5 de servicios concentra el 92.66% del impacto global de la organización:", bold_prefix="Concentración del Esfuerzo en el Top 5: ")
add_bullet_point(doc, "1. S-04 Aula Virtual: Brecha ponderada de 129 600 puntos (33.13% del esfuerzo total).")
add_bullet_point(doc, "2. S-03 Pago de Pensiones y Recaudación: Brecha ponderada de 97 200 puntos (24.85% del esfuerzo).")
add_bullet_point(doc, "3. S-02 Registro y Publicación de Actas: Brecha ponderada de 90 720 puntos (23.19% del esfuerzo).")
add_bullet_point(doc, "4. S-09 Préstamo de Biblioteca y OPAC: Brecha ponderada de 25 500 puntos (6.52% del esfuerzo).")
add_bullet_point(doc, "5. S-01 Matrícula Semestral: Brecha ponderada de 19 440 puntos (4.97% del esfuerzo).")

add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "graficos", "GD_servicios.png"),
                         "Comparativo del nivel de digitalización actual frente al objetivo del PETI por servicio. Fuente: Elaboración propia mediante GD02_analisis_servicios.py (2026).", 1)

doc.add_page_break()

# -------------------------------------------------------------------------
# PASO B
# -------------------------------------------------------------------------
add_styled_heading(doc, "2.2 Paso B — Diagnóstico de los ocho ejes con línea base", level=2)
add_styled_paragraph(doc, "Se evaluaron los ocho ejes de la arquitectura de gobierno digital peruana (RSGD 005-2018-PCM/SEGDI y D. S. 029-2021-PCM) adaptados a una institución educativa privada. Se formularon 23 indicadores técnicos y de gestión, asignando a cada uno su línea base actual, fuente de verificación, evidencia y nivel de madurez (1 a 5).")

add_styled_paragraph(doc, "El archivo generado en el repositorio es 06_gobierno_digital/GD03_ejes.csv. A continuación se presenta la matriz diagnóstica completa:")

# Tabla de Ejes
ejes_headers = ["Eje", "Indicador Clave", "Línea Base", "Fuente y Fecha", "Nivel", "Observación Técnica"]
ejes_rows = [
    ejes_headers,
    ["1. Gobernanza", "¿Comité de TI con actas y reuniones?", "No (Existe en papel, 0 reuniones)", "Entrevista Dirección, jun 2026", "1", "Inexistencia de liderazgo corporativo para priorizar."],
    ["1. Gobernanza", "¿Plan de TI (PETI) vigente?", "No", "Revisión documental, jun 2026", "1", "Este PETI será el primero formulado en 16 años."],
    ["1. Gobernanza", "% presupuesto asignado a TI", "1.81 % (S/ 356K de S/ 19.7M)", "Presupuesto TI 2025", "2", "Muy inferior al estándar del sector privado (3.5%-5%)."],
    ["2. Servicios", "Nivel digitalización ponderado", "2.31 de 5.00", "Script GD02 / Catálogo GD01", "2", "Distorsionado por lecturas Moodle; ventanilla atrasada."],
    ["2. Servicios", "% trámites 100% digitales", "0 %", "Auditoría mesa de partes", "1", "Todos los trámites exigen firma física o presencialidad."],
    ["2. Servicios", "Satisfacción con servicios TI", "No se mide", "Cuestionario institucional", "1", "Primer proyecto: Implementar CSAT/NPS en mesa de ayuda."],
    ["3. Interoperab.", "N.º integraciones automáticas", "0 (Sistemas aislados)", "Inspección de arquitectura", "1", "Transferencia manual de archivos planos entre sistemas."],
    ["3. Interoperab.", "Interoperabilidad con Estado", "0 integraciones (Carga manual)", "Secretaría Académica", "1", "Sin conexión web service con PIDE, RENIEC ni SUNAT."],
    ["3. Interoperab.", "Documentos redundantes pedidos", "65.38 % (17 de 26)", "Prueba 'una sola vez' GD02", "1", "Severa ineficiencia y fricción administrativa."],
    ["4. Seguridad", "¿Existe SGSI (ISO 27001)?", "No existe", "Diagnóstico de cumplimiento", "1", "Exposición legal crítica ante la ANPD (Ley 29733)."],
    ["4. Seguridad", "Incidentes reportados formalmente", "0 (1 crítico admitido: 17 notas)", "Ficha de hechos Sección G", "1", "Subregistro total; 17 notas alteradas impunemente."],
    ["4. Seguridad", "% sistemas con backup probado", "0 %", "Ficha TI / Entrevistas soporte", "1", "Moodle sin respaldos; 3 sistemas sin copia verificada."],
    ["5. Datos", "Entidades maestras con custodio", "0 de 7 entidades", "Revisión mesa de ayuda", "1", "Notas, matrícula y pagos huérfanos de dueño formal."],
    ["5. Datos", "% decisiones con datos analíticos", "No se mide", "Mesa de servicio sección 6", "1", "Primer proyecto: Data Warehouse Institucional y BI."],
    ["5. Datos", "Cumplimiento Ley 29733", "No implementado", "Inventario regulatorio", "1", "Tratamiento de datos de menores sin cláusulas legales."],
    ["6. Identidad", "¿Directorio único institucional (SSO)?", "No", "Inventario de accesos", "1", "Credenciales disjuntas y desarticuladas."],
    ["6. Identidad", "% sistemas con 2FA (MFA)", "20 % (1 de 5: correo en nube)", "Auditoría técnica de accesos", "2", "Sistema de notas y Moodle operan con password simple."],
    ["6. Identidad", "Políticas técnicas de passwords", "No exigible técnicamente", "Ficha TI / Caso notas", "1", "Sistema legado no soporta expiración ni complejidad."],
    ["7. Talento", "% personal con capacitación TI", "0 % formal", "RRHH / Cuestionario P6", "1", "Solo correos de concientización sin confirmación."],
    ["7. Talento", "Ratio soporte TI vs usuarios", "1 técnico por 1 146 usuarios", "Planilla TI (3 p.) / 3 440", "1", "Déficit crítico que colapsa la atención operativa."],
    ["8. Infraestruc.", "Disponibilidad de sistemas (Uptime)", "No se mide", "Mesa servicio: 22 caídas/año", "1", "Primer proyecto: Monitoreo Zabbix y SLAs 24/7."],
    ["8. Infraestruc.", "% equipos con soporte vigente", "40 %", "Inventario de hardware", "2", "Servidor principal con 9 años de antigüedad al límite."],
    ["8. Infraestruc.", "Monitoreo de enlaces WAN/QoS", "No se mide", "Mesa servicio: 11 caídas/año", "1", "Primer proyecto: Telemetría de red y control de QoS."]
]
tbl_ejes = doc.add_table(rows=len(ejes_rows), cols=6)
for r_i, row in enumerate(ejes_rows):
    for c_i, val in enumerate(row):
        tbl_ejes.rows[r_i].cells[c_i].paragraphs[0].text = val
style_table(tbl_ejes)

add_styled_paragraph(doc, "Para consolidar los resultados se ejecutó el script 06_gobierno_digital/GD04_madurez.py, obteniendo la siguiente salida:")

with open(os.path.join(SCRIPT_DIR, "evidencias", "S11", "salidas", "GD04_salida.txt"), "r", encoding="utf-8") as f:
    gd04_txt = f.read()

p_log2 = doc.add_paragraph()
p_log2.paragraph_format.left_indent = Inches(0.2)
p_log2.paragraph_format.right_indent = Inches(0.2)
p_log2.paragraph_format.space_before = Pt(4)
p_log2.paragraph_format.space_after = Pt(6)
r_log2 = p_log2.add_run(gd04_txt)
r_log2.font.name = "Consolas"
r_log2.font.size = Pt(7.5)
r_log2.font.color.rgb = RGBColor(30, 41, 59)

add_styled_paragraph(doc, "La evaluación integral de los 23 indicadores arroja un puntaje promedio institucional de 1.17 sobre 5.00, ubicando a la entidad en el NIVEL 1: INICIAL / AD-HOC (Vulnerable y Reactivo). Las prácticas tecnológicas dependen de la buena voluntad individual, carecen de controles automatizados y exhiben brechas críticas en ciberseguridad y continuidad.", bold_prefix="Nivel de Madurez y Posicionamiento: ")

add_styled_paragraph(doc, "Se detectaron cuatro (4) indicadores en estado 'No se mide'. La metodología de planeamiento estratégico estipula que sin línea base es imposible comprometer metas. Por tanto, el PETI establece los siguientes cuatro proyectos de instrumentación obligatorios e inmediatos:", bold_prefix="Tratamiento Metodológico de Indicadores Sin Línea Base: ")
add_bullet_point(doc, "Satisfacción de usuarios con servicios TI -> Proyecto INST-01: Mesa de Ayuda Omnicanal con Evaluación Automatizada CSAT/NPS.")
add_bullet_point(doc, "% decisiones basadas en datos analíticos -> Proyecto INST-02: Almacén de Datos Institucional y Cuadro de Mando Integral.")
add_bullet_point(doc, "Disponibilidad efectiva de sistemas críticos (Uptime) -> Proyecto INST-03: Despliegue de Sistema de Monitoreo Proactivo Zabbix y SLAs.")
add_bullet_point(doc, "Monitoreo proactivo de capacidad y ancho de banda WAN -> Proyecto INST-04: Instrumentación de Telemetría SNMP/NetFlow y QoS.")

add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "graficos", "GD_madurez.png"),
                         "Radar de madurez de gobierno digital por eje (Nivel actual 1.17 vs. Meta PETI 3.50). Fuente: Elaboración propia mediante GD04_madurez.py (2026).", 2)

doc.add_page_break()

# -------------------------------------------------------------------------
# PASO C
# -------------------------------------------------------------------------
add_styled_heading(doc, "2.3 Paso C — Despliegue del tablero de línea base en Metabase", level=2)
add_styled_paragraph(doc, "Conforme a la metodología, se desplegó un entorno analítico en local mediante Docker con PostgreSQL 16 (puerto 5433) y Metabase (puerto 3001). Mediante el script 06_gobierno_digital/GD05_carga_tablero.py se migraron y normalizaron las cuatro matrices analíticas del diagnóstico:")
add_bullet_point(doc, "Tabla 'servicios': Catálogo de los 10 servicios con volumen, canales y niveles de digitalización (GD01).")
add_bullet_point(doc, "Tabla 'ejes_gd': Diagnóstico de los 23 indicadores de los ocho ejes con líneas base y observaciones (GD03).")
add_bullet_point(doc, "Tabla 'cumplimiento': Matriz de cumplimiento legal frente a MINEDU, SUNEDU, ANPD (Ley 29733) y SUNAT (NO03).")
add_bullet_point(doc, "Tabla 'capacidad': Matriz de capacidad de procesos de TI bajo el marco COBIT 2019 (MG03).")

with open(os.path.join(SCRIPT_DIR, "evidencias", "S11", "salidas", "GD05_salida.txt"), "r", encoding="utf-8") as f:
    gd05_txt = f.read()

p_log3 = doc.add_paragraph()
p_log3.paragraph_format.left_indent = Inches(0.2)
p_log3.paragraph_format.right_indent = Inches(0.2)
p_log3.paragraph_format.space_before = Pt(4)
p_log3.paragraph_format.space_after = Pt(6)
r_log3 = p_log3.add_run(gd05_txt)
r_log3.font.name = "Consolas"
r_log3.font.size = Pt(7.5)
r_log3.font.color.rgb = RGBColor(30, 41, 59)

add_styled_paragraph(doc, "En Metabase se construyó el Tablero Maestro de Línea Base estructurado en seis visualizaciones esenciales que responden a las interrogantes estratégicas del gobierno corporativo:")

add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "graficos", "tablero_metabase_completo.png"),
                         "Vista consolidada del Tablero de Línea Base en Metabase desplegado para el Instituto Superior Tecnológico Vertiente Sur. Fuente: Metabase BI sobre PostgreSQL (2026).", 3, width_in=6.4)

add_styled_heading(doc, "Detalle de las Seis Visualizaciones Mínimas del Tablero", level=3)

# Tarjeta 1
add_styled_paragraph(doc, "Responde a la pregunta: ¿Qué servicios están más atrasados? Muestra que los servicios de admisión, titulación, constancias, matrículas y cobranzas están en Nivel 1, mientras que el soporte de TI está en Nivel 0.", bold_prefix="Visualización 1 — Nivel de digitalización por servicio (Barras Horizontales): ")
add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "evidencias", "S11", "capturas", "01_nivel_digitalizacion_servicios.png"),
                         "Visualización 1: Comparativa del nivel actual frente a la meta PETI por servicio. Fuente: Metabase (2026).", 4, width_in=5.2)

# Tarjeta 2
add_styled_paragraph(doc, "Responde a la pregunta: ¿Qué proporción de las atenciones es realmente digital? Revela un valor de 2.31 sobre 5.00, contrastado con el promedio simple de 1.20, demostrando la concentración del volumen en el Moodle y la precariedad de los trámites de ventanilla.", bold_prefix="Visualización 2 — Nivel ponderado por volumen (Número Grande / KPI): ")
add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "evidencias", "S11", "capturas", "02_nivel_ponderado_kpi.png"),
                         "Visualización 2: Tarjeta KPI de nivel medio ponderado por volumen transaccional. Fuente: Metabase (2026).", 5, width_in=4.2)

# Tarjeta 3
add_styled_paragraph(doc, "Responde a la pregunta: ¿Dónde está la mayor brecha de gobierno digital? Identifica que Interoperabilidad, Seguridad, Datos y Talento se encuentran en el nivel basal crítico de 1.00.", bold_prefix="Visualización 3 — Madurez por eje de gobierno digital (Barras Horizontales): ")
add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "evidencias", "S11", "capturas", "03_madurez_ejes_barras.png"),
                         "Visualización 3: Nivel de madurez alcanzado en los 8 ejes del gobierno digital. Fuente: Metabase (2026).", 6, width_in=5.2)

# Tarjeta 4
add_styled_paragraph(doc, "Responde a la pregunta: ¿Cuántas obligaciones se incumplen? Evidencia que el 43% de las normas (3 de 7) no se cumplen en absoluto (principalmente Ley 29733 de Datos Personales), el 43% se cumple parcialmente y solo el 14% (facturación electrónica) cumple con la normativa.", bold_prefix="Visualización 4 — Estado de cumplimiento normativo (Gráfico de Anillo): ")
add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "evidencias", "S11", "capturas", "04_cumplimiento_normativo_anillo.png"),
                         "Visualización 4: Proporción de obligaciones legales cumplidas, parciales e incumplidas. Fuente: Metabase (2026).", 7, width_in=4.2)

# Tarjeta 5
add_styled_paragraph(doc, "Responde a la pregunta: ¿Qué procesos de TI exigen mayor esfuerzo de transformación? Los procesos de gestión de seguridad (APO13), gestión de problemas (DSS03) y evaluación de desempeño (MEA01) parten de nivel 0, requiriendo un salto cuantitativo de 3 niveles.", bold_prefix="Visualización 5 — Capacidad de procesos COBIT 2019 (Barras Agrupadas): ")
add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "evidencias", "S11", "capturas", "05_capacidad_procesos_agrupadas.png"),
                         "Visualización 5: Capacidad actual vs. capacidad objetivo por proceso COBIT 2019. Fuente: Metabase (2026).", 8, width_in=5.2)

# Tarjeta 6
add_styled_paragraph(doc, "Responde a la pregunta: ¿Qué necesitamos empezar a medir con urgencia? Sintetiza los cuatro indicadores ciegos del instituto y formaliza su proyecto de instrumentación en el PETI.", bold_prefix="Visualización 6 — Indicadores sin línea base (Tabla Ejecutiva): ")
add_figure_with_caption(doc, os.path.join(SCRIPT_DIR, "evidencias", "S11", "capturas", "06_indicadores_sin_linea_base_tabla.png"),
                         "Visualización 6: Tabla analítica de indicadores sin línea base y proyectos asignados. Fuente: Metabase (2026).", 9, width_in=5.6)

doc.add_page_break()

# -------------------------------------------------------------------------
# PASO D
# -------------------------------------------------------------------------
add_styled_heading(doc, "2.4 Paso D — Redacción de la Sección 6.1 del PETI", level=2)
add_styled_paragraph(doc, "Conforme a los estándares de planeamiento estratégico de TI, se redactó íntegramente la Sección 6.1 del PETI en el archivo 06_gobierno_digital/6.1_situacion_actual.md. A continuación se presentan los apartados sustantivos articulados:")

add_styled_paragraph(doc, "La Sección 6.1.1 fundamenta cómo los componentes del gobierno digital peruano aplican a la entidad privada en función de su supervisión por MINEDU (Condiciones Básicas de Calidad), SUNEDU y fiscalización de la ANPD por tratamiento de bancos de datos personales de estudiantes y menores de edad.", bold_prefix="6.1.1 Marco Conceptual y Alcance: ")

add_styled_paragraph(doc, "La Sección 6.1.2 incorpora el catálogo auditado y sustenta cuantitativamente por qué el promedio ponderado (2.31) desenmascara la ineficiencia de los trámites troncales de ventanilla que operan en Nivel 1 y 0.", bold_prefix="6.1.2 Catálogo de Servicios y Digitalización: ")

add_styled_paragraph(doc, "La Sección 6.1.3 demuestra que el 65.38% de los requisitos documentales solicitados ya están en poder del instituto, disponiendo su erradicación técnica mediante interoperabilidad.", bold_prefix="6.1.3 Principio de 'Una Sola Vez': ")

add_styled_paragraph(doc, "La Sección 6.1.4 consolida la matriz de 23 indicadores evaluados sobre los ocho ejes con fuentes, evidencias y niveles.", bold_prefix="6.1.4 Diagnóstico por Eje: ")

add_styled_paragraph(doc, "La Sección 6.1.5 formaliza los proyectos INST-01 a INST-04 para dotar de línea base a los cuatro procesos desprovistos de telemetría.", bold_prefix="6.1.5 Indicadores Sin Línea Base: ")

add_styled_paragraph(doc, "La Sección 6.1.6 ubica a la entidad en el Nivel 1 de Madurez (Inicial / Ad-hoc) sustentando la necesidad de formalizar la gestión de servicios y la seguridad criptográfica.", bold_prefix="6.1.6 Nivel de Madurez de Gobierno Digital: ")

add_styled_paragraph(doc, "La Sección 6.1.7 sintetiza las brechas críticas: el 92.66% del esfuerzo se focaliza en 5 servicios troncales y los ejes basales de Seguridad, Interoperabilidad, Datos y Talento.", bold_prefix="6.1.7 Síntesis de Brechas de Mayor Impacto: ")

# -------------------------------------------------------------------------
# PASO E
# -------------------------------------------------------------------------
add_styled_heading(doc, "2.5 Paso E — Validación cruzada y comprobaciones de integridad", level=2)
add_styled_paragraph(doc, "Para garantizar la autenticidad y solidez técnica del diagnóstico, se ejecutaron las tres comprobaciones obligatorias estipuladas en la metodología:")

add_styled_paragraph(doc, "Comprobación 1 — Recorrido completo de dos servicios auditados:", bold_prefix="Validación 1: ")
add_bullet_point(doc, "Servicio S-01 (Matrícula): Se recorrió el flujo comprobando que el nivel declarado (1) es el real: aunque existe una pantalla preliminar en intranet, el alumno se desplaza físicamente al banco, solicita el sello manual de tesorería y hace fila en secretaría académica para entregar 4 papeles.")
add_bullet_point(doc, "Servicio S-02 (Registro de Actas): Se auditó el flujo confirmando el Nivel 2: el docente digita notas pero debe imprimir y firmar el acta físicamente. La carencia de 2FA y la debilidad de contraseñas permitió que durante la licencia de un docente se alteraran fraudulentamente 17 notas.")

add_styled_paragraph(doc, "Comprobación 2 — Respaldo de evidencias en los ocho ejes:", bold_prefix="Validación 2: ")
add_styled_paragraph(doc, "Se verificó que los puntajes basales (1) no fueron asignados subjetivamente, sino contrastados con hechos verificados: ausencia de actas del Comité de TI, ausencia de SGSI formal, registros de mesa de ayuda con 125 tickets abiertos crónicos, subregistro de incidentes y falta de pruebas de restauración de respaldos.")

add_styled_paragraph(doc, "Comprobación 3 — Auditoría física del principio de 'una sola vez':", bold_prefix="Validación 3: ")
add_styled_paragraph(doc, "Se comprobó documentalmente sobre requisitos tangibles: el alumno entrega copia de DNI ya registrado al ingresar, ficha emitida por el software institucional y voucher de un pago que ya figura en el módulo de cobranzas.")

# -------------------------------------------------------------------------
# PASO F
# -------------------------------------------------------------------------
add_styled_heading(doc, "2.6 Paso F — Control de versiones y transferencia técnica", level=2)
add_styled_paragraph(doc, "Todos los artefactos, scripts, bases de datos y matrices fueron debidamente versionados en el repositorio oficial de GitHub del equipo:")
add_styled_paragraph(doc, "https://github.com/mauchoquenachoque/PETI-SEMANA11-CHOQUENA-LAOS-RAMOSL-CALLOTICONA", bold_prefix="URL Oficial del Repositorio: ")
add_styled_paragraph(doc, "Se crearon las etiquetas git correspondientes: 'v0.11' (versión de avance del PETI) y 'taller-11' (etiqueta de entrega evaluativa).", bold_prefix="Etiquetas de Versión: ")

add_styled_heading(doc, "Pregunta de Transferencia Técnica", level=3)
add_styled_paragraph(doc, "¿Qué riesgo correría una organización real si este diagnóstico y línea base se hicieran mal?", bold_prefix="Interrogante: ")
add_styled_paragraph(doc, "Respuesta Técnica del Equipo:", bold_prefix="Dictamen Pericial: ")
add_styled_paragraph(doc, "Si una institución educativa real formulara este diagnóstico de forma superficial o autocomplaciente —guiándose únicamente por el promedio simple o declarando niveles de digitalización teóricos sin contrastar la ventanilla física— incurriría en cuatro riesgos catastróficos:")
add_bullet_point(doc, "1. Riesgo de Fraude Académico e Impunidad Legal: Al no reconocer la ausencia de identidad digital única, auditoría criptográfica y segundo factor (MFA), el sistema permanece expuesto a la adulteración ilícita de actas y venta de certificados falsos, lo que acarrea la cancelación del licenciamiento institucional por parte del MINEDU/SUNEDU y denuncias penales contra la Dirección.")
add_bullet_point(doc, "2. Riesgo Sancionador Severo ante la ANPD: Declarar cumplimiento en el eje de datos sin auditar el tratamiento diferenciado de datos de menores ni inscribir los bancos de datos personales expone a la institución a multas coercitivas de hasta 100 UIT conforme a la Ley 29733.")
add_bullet_point(doc, "3. Riesgo de Inversión Tecnológica Ciega: Asignar presupuesto a proyectos sobre indicadores en estado 'No se mide' equivale a ejecutar gasto sin capacidad de medir retorno ni impacto, perpetuando silos tecnológicos que no resuelven la frustración del estudiante.")
add_bullet_point(doc, "4. Colapso Operacional en Períodos Críticos: Omitir la medición de disponibilidad (Uptime) y capacidad de servidores provoca que en cada cierre de semestre el sistema colapse (22 incidentes anuales reconocidos), paralizando la recaudación financiera y dañando severamente la reputación de la institución.")

doc.add_page_break()

# =========================================================================
# 3. RESULTADOS
# =========================================================================
add_styled_heading(doc, "3. Resultados y evidencias procedimentales", level=1)

add_styled_heading(doc, "3.1 Los tres resultados que se califican", level=2)
add_styled_paragraph(doc, "Conforme a la rúbrica de evaluación procedimental, se sustentan los tres entregables troncales del taller:")
add_bullet_point(doc, "El Catálogo con Niveles Comprobados: Catálogo exhaustivo de 10 servicios con volumen transaccional auditado (231 309 operaciones), tiempos de respuesta, canales y justificación de metas. Evidencia en 06_gobierno_digital/GD01_catalogo_servicios.csv.", "1. Catálogo de Servicios: ")
add_bullet_point(doc, "Los Ocho Ejes con Evidencia: Diagnóstico de 23 indicadores con fuentes fechadas, evidencias tangibles y niveles de madurez asignados con rigor. Evidencia en 06_gobierno_digital/GD03_ejes.csv.", "2. Ocho Ejes Evaluados: ")
add_bullet_point(doc, "La Línea Base del Plan y Metabase: Tablero interactivo desplegado en Metabase con las seis visualizaciones mínimas y formalización de proyectos de instrumentación para los indicadores 'No se mide'. Evidencia en graficos/tablero_metabase_completo.png y evidencias/S11/capturas/.", "3. Línea Base y Tablero: ")

add_styled_heading(doc, "3.2 Lista de comprobación del taller (13 Verificaciones)", level=2)
add_styled_paragraph(doc, "A continuación se presenta el cumplimiento exhaustivo de la lista de comprobación de la sección 3.2 de la guía:")

# Tabla 1 de Comprobación
chk_headers = ["#", "Resultado Esperado", "¿Se logró?", "Evidencia en el Repositorio GitHub"]
chk_rows = [
    chk_headers,
    ["1", "Catálogo con ≥ 7 servicios con volumen y nivel", "Logrado", "06_gobierno_digital/GD01_catalogo_servicios.csv (10 servicios catalogados)"],
    ["2", "Nivel medio ponderado vs. promedio simple", "Logrado", "Salida de GD02_analisis_servicios.py (Simple: 1.20 | Ponderado: 2.31)"],
    ["3", "Prueba del principio de una sola vez", "Logrado", "Salida GD02: 17 de 26 documentos redundantes (65.38% de redundancia)"],
    ["4", "Brecha por servicio ponderada (Top 5 impacto)", "Logrado", "Salida GD02: Top 5 concentra el 92.66% del esfuerzo de transformación"],
    ["5", "Gráfico del catálogo actual vs. objetivo", "Logrado", "graficos/GD_servicios.png insertado en el informe"],
    ["6", "Ocho ejes evaluados con ≥ 2 indicadores c/u", "Logrado", "06_gobierno_digital/GD03_ejes.csv (23 indicadores evaluados)"],
    ["7", "Línea base declarada o 'no se mide' explícito", "Logrado", "Columna 'Línea base' en GD03_ejes.csv con fuentes fechadas"],
    ["8", "Indicadores sin línea base contados con proyecto", "Logrado", "Salida GD04: 4 indicadores identificados con proyectos INST-01 a INST-04"],
    ["9", "Nivel de madurez global y posicionamiento", "Logrado", "Salida GD04: 1.17 de 5.00 -> Nivel 1 Inicial / Ad-hoc"],
    ["10", "Radar de madurez por eje", "Logrado", "graficos/GD_madurez.png insertado en el informe"],
    ["11", "Tablero en Metabase con ≥ 6 visualizaciones", "Logrado", "graficos/tablero_metabase_completo.png y 6 capturas en evidencias/S11/capturas/"],
    ["12", "Sección 6.1 redactada con síntesis de brechas", "Logrado", "06_gobierno_digital/6.1_situacion_actual.md articulado"],
    ["13", "Etiqueta v0.11 y taller-11 en Git", "Logrado", "Etiquetas v0.11 y taller-11 creadas y versionadas en GitHub"]
]
tbl_chk = doc.add_table(rows=len(chk_rows), cols=4)
for r_i, row in enumerate(chk_rows):
    for c_i, val in enumerate(row):
        tbl_chk.rows[r_i].cells[c_i].paragraphs[0].text = val
style_table(tbl_chk)

add_styled_heading(doc, "3.3 Evaluación de la rúbrica procedimental (20 puntos)", level=2)
add_styled_paragraph(doc, "El equipo autoevalúa el desempeño procedimental alcanzado en función de los criterios de la guía del taller:")
add_bullet_point(doc, "El Criterio de Éxito (4/4 puntos): Se cumplió al 100% el reto de la sesión, articulando el caso simulado de educación superior tecnológica con herramientas reales de analítica.", "• Criterio 1: ")
add_bullet_point(doc, "El Catálogo con Niveles (4/4 puntos): Levantamiento completo y riguroso de 10 servicios con volumen, nivel comprobado en ventanilla y justificación técnica.", "• Criterio 2: ")
add_bullet_point(doc, "Los Ocho Ejes con Evidencia (4/4 puntos): Evaluación de 23 indicadores soportados en hechos verificables, fechas de corte y fuentes documentales.", "• Criterio 3: ")
add_bullet_point(doc, "La Línea Base del Plan (4/4 puntos): Cifras de partida reales, delimitación de indicadores 'No se mide' y asignación de proyectos de instrumentación.", "• Criterio 4: ")
add_bullet_point(doc, "La Validación (4/4 puntos): Ejecución de las tres comprobaciones cruzadas, documentando con honestidad técnica las deficiencias detectadas y las acciones de mitigación.", "• Criterio 5: ")
add_styled_paragraph(doc, "Puntaje Total Obtenido: 20 / 20 puntos (Nivel Logrado en todas las dimensiones).", bold_prefix="Calificación Global: ")

add_styled_heading(doc, "3.4 Problemas identificados y oportunidades de mejora", level=2)
add_styled_paragraph(doc, "1. Problema: El sistema académico legado opera sobre un servidor local de 2017 que no permite técnicamente exigir políticas de complejidad ni expiración de contraseñas, lo que derivó en la alteración ilícita de 17 actas de notas.")
add_styled_paragraph(doc, "   -> Mejora implementada en PETI: Se formuló el Proyecto SEG-01 para la incorporación inmediata de un middleware de autenticación federada con Segundo Factor (MFA/OTP) y firma digital con sellado de tiempo.")
add_styled_paragraph(doc, "2. Problema: La mesa de ayuda operaba mediante una hoja de cálculo no estructurada que un practicante llenaba al final de la jornada, generando un subregistro superior al 50% y 125 tickets abiertos crónicos.")
add_styled_paragraph(doc, "   -> Mejora implementada en PETI: Se estructuró el Proyecto INST-01 para desplegar una plataforma ITSM (GLPI/Jira) con ANS formalizados y catalogación automática de incidentes y requerimientos.")

doc.add_page_break()

# =========================================================================
# 4. CONCLUSIONES
# =========================================================================
add_styled_heading(doc, "4. Conclusiones estratégicas y técnicas", level=1)
add_styled_paragraph(doc, "1. El nivel de digitalización ponderado por volumen transaccional es la métrica rectora para el planeamiento de TI. Evaluar la organización únicamente por el promedio aritmético simple sobrevalora el avance cuando los servicios masivos son analógicos, o genera un falso confort cuando una plataforma secundaria (como foros de Moodle) concentra tráfico pasivo mientras los servicios misionales de ventanilla colapsan.", bold_prefix="Conclusión 1 (Métrica Ponderada): ")
add_styled_paragraph(doc, "2. Un indicador sin línea base no puede sostener una meta estratégica ni demostrar avance ante el gobierno corporativo. Declarar metas porcentuales arbitrarias sobre variables no medidas constituye una mala práctica de planeamiento; el primer compromiso del PETI ante indicadores ciegos no es prometer mejoras, sino ejecutar proyectos de instrumentación técnica y telemetría.", bold_prefix="Conclusión 2 (Instrumentación de Línea Base): ")
add_styled_paragraph(doc, "3. Gobierno digital no equivale a la proliferación de sistemas de información aislados, sino a la capacidad del usuario de completar una transacción integral sin fricción ni cambio forzado de canal. La presencia de cinco sistemas desconectados en el instituto generó un índice de redundancia documental del 65.38%, demostrando que la falta de interoperabilidad traslada el costo del desorden institucional al usuario.", bold_prefix="Conclusión 3 (Interoperabilidad vs. Silos): ")
add_styled_paragraph(doc, "4. La seguridad de la información y la identidad digital son los habilitadores indispensables de la confianza institucional. El incidente de 17 actas alteradas evidenció que la digitalización sin controles criptográficos, firma digital y trazabilidad inmutable no genera modernización, sino un grave riesgo reputacional y legal frente a las entidades fiscalizadoras.", bold_prefix="Conclusión 4 (Seguridad como Habilitador): ")

doc.add_page_break()

# =========================================================================
# 5. CUESTIONARIO
# =========================================================================
add_styled_heading(doc, "5. Cuestionario de evaluación técnica", level=1)

add_styled_heading(doc, "Pregunta 1", level=2)
add_styled_paragraph(doc, "¿Cuál es la diferencia metodológica entre el promedio simple y el promedio ponderado por volumen en el catálogo de servicios, y por qué el segundo es determinante para la toma de decisiones?", bold_prefix="Enunciado: ")
add_styled_paragraph(doc, "El promedio simple asigna el mismo peso relativo a un trámite marginal (ej. trámite de titulación con 480 solicitudes al año) que a un servicio masivo (ej. acceso al aula virtual con 129 600 interacciones o pago de pensiones con 32 400 operaciones). Esto genera una severa distorsión estadística: si diez trámites de bajo volumen son digitales pero los dos servicios de mayor demanda son totalmente analógicos, el promedio simple mostrará un avance ficticio. El promedio ponderado multiplica el nivel de cada servicio por su frecuencia anual de ocurrencia, midiendo la verdadera experiencia digital de la masa usuaria y priorizando inversiones donde el retorno operativo y la erradicación de colas sean máximos.", bold_prefix="Respuesta Fundamentada: ")

add_styled_heading(doc, "Pregunta 2", level=2)
add_styled_paragraph(doc, "¿En qué consiste el principio de 'una sola vez' (Once-Only Principle) y cuáles son las repercusiones de su transgresión en una institución educativa?", bold_prefix="Enunciado: ")
add_styled_paragraph(doc, "Establecido en la Ley de Gobierno Digital (D.L. 1412) y estándares de la OCDE, estipula que la administración no debe exigir al administrado documentos o datos que ya consten en sus archivos o que hayan sido expedidos por sus propias unidades orgánicas. En el Instituto Vertiente Sur, el índice de redundancia alcanzó el 65.38% (17 de 26 documentos). La repercusión es directa: saturación de ventanillas, colas de más de dos horas en matrícula y emisión de certificados, gasto superfluo en impresiones, demoras de hasta 10 días para trámites que deberían resolverse en segundos e incentivo a la corrupción o falsificación de requisitos físicos.", bold_prefix="Respuesta Fundamentada: ")

add_styled_heading(doc, "Pregunta 3", level=2)
add_styled_paragraph(doc, "¿Por qué se establece como axioma que 'sin línea base no puede fijarse meta ni demostrarse avance'?", bold_prefix="Enunciado: ")
add_styled_paragraph(doc, "Una meta estratégica en un PETI representa una tasa de variación cuantitativa respecto a un estado de partida verificable. Si un indicador se encuentra en estado 'No se mide' (como la disponibilidad de servidores o la satisfacción de usuarios en el caso analizado), cualquier meta fijada carece de sustento empírico y credibilidad financiera. No se puede calcular el esfuerzo requerido ni demostrar si una inversión de S/ 100 000 produjo una mejora del 10% o un retroceso. Por ende, la disciplina de planeamiento exige que el primer paso sea instrumentar la medición mediante sondas y software de monitoreo.", bold_prefix="Respuesta Fundamentada: ")

add_styled_heading(doc, "Pregunta 4", level=2)
add_styled_paragraph(doc, "¿Cómo se adaptan los ocho ejes de gobierno digital a una organización educativa privada regulada?", bold_prefix="Enunciado: ")
add_styled_paragraph(doc, "Aunque la organización es privada, los ocho ejes son perfectamente homologables: la Gobernanza se materializa en el Comité de TI y el PETI; los Servicios en la matrícula y pagos online; la Interoperabilidad en la comunicación sin fricción entre Sistema Académico, Moodle y reportes a MINEDU/SUNEDU; la Seguridad Digital en el SGSI para evitar fraude en notas; los Datos en la gobernanza y protección de datos personales de menores (Ley 29733); la Identidad Digital en cuentas SSO corporativas y MFA; el Talento en competencias digitales de docentes; y la Infraestructura en la alta disponibilidad y redundancia de servidores y enlaces.", bold_prefix="Respuesta Fundamentada: ")

add_styled_heading(doc, "Pregunta 5", level=2)
add_styled_paragraph(doc, "¿Qué valor estratégico aporta consolidar el diagnóstico en un tablero de BI como Metabase sobre PostgreSQL frente a mantenerlo en hojas de cálculo aisladas?", bold_prefix="Enunciado: ")
add_styled_paragraph(doc, "El diagnóstico en hojas de cálculo adolece de inconsistencia, desactualización y riesgo de adulteración (el caso demostró discrepancias en cifras entre áreas y 125 tickets abiertos crónicos). Desplegar el tablero en Metabase sobre un modelo relacional en PostgreSQL garantiza integridad referencial, actualización en tiempo real, trazabilidad de consultas SQL, gobernanza del dato y democratización visual para los directores y comités, transformando el diagnóstico estático en un instrumento continuo de supervisión para la Sección 10 del PETI.", bold_prefix="Respuesta Fundamentada: ")

doc.add_page_break()

# =========================================================================
# 6. REFERENCIAS BIBLIOGRÁFICAS
# =========================================================================
add_styled_heading(doc, "6. Referencias bibliográficas (Normas APA)", level=1)
referencias = [
    "Decreto Legislativo N.° 1412. (2018, 13 de setiembre). Ley de Gobierno Digital. Diario Oficial El Peruano. https://www.gob.pe/institucion/pcm/colecciones/147-normativa-sobre-gobierno-digital",
    "Decreto Supremo N.° 029-2021-PCM. (2021, 19 de febrero). Reglamento del Decreto Legislativo N.° 1412, Ley de Gobierno Digital. Presidencia del Consejo de Ministros. https://www.gob.pe/13326-reglamento-de-la-ley-de-gobierno-digital",
    "Decreto Supremo N.° 085-2023-PCM. (2023, 26 de julio). Política Nacional de Transformación Digital al 2030. Diario Oficial El Peruano. https://busquedas.elperuano.pe/dispositivo/NL/2200457-5",
    "ISACA. (2019). COBIT 2019 Framework: Governance and Management Objectives. Information Systems Audit and Control Association.",
    "ISO/IEC. (2018). ISO/IEC 20000-1:2018 Information technology — Service management — Part 1: Service management system requirements. International Organization for Standardization.",
    "ISO/IEC. (2022). ISO/IEC 27001:2022 Information security, cybersecurity and privacy protection — Information security management systems — Requirements. International Organization for Standardization.",
    "Ley N.° 29733. (2011, 3 de julio). Ley de Protección de Datos Personales. Diario Oficial El Peruano.",
    "Metabase. (2026). Metabase Documentation: Business Intelligence and Analytics Open Source. https://www.metabase.com/docs/latest/",
    "OECD. (2020). The OECD Digital Government Policy Framework: Six dimensions of a digital government. OECD Public Governance Policy Papers, No. 02, OECD Publishing. https://doi.org/10.1787/f64fed2a-en",
    "Resolución de Secretaría de Gobierno Digital N.° 005-2018-PCM/SEGDI. (2018, 29 de noviembre). Lineamientos para la formulación del Plan de Gobierno Digital. Presidencia del Consejo de Ministros."
]
for ref in referencias:
    p_r = doc.add_paragraph()
    p_r.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_r.paragraph_format.left_indent = Inches(0.4)
    p_r.paragraph_format.first_line_indent = Inches(-0.4)
    p_r.paragraph_format.space_after = Pt(6)
    p_r.paragraph_format.line_spacing = 1.15
    r = p_r.add_run(ref)
    r.font.name = "Calibri"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(30, 41, 59)

doc.add_page_break()

# =========================================================================
# 7. ANEXOS
# =========================================================================
add_styled_heading(doc, "7. Anexos documentales, analíticos y gráficos", level=1)
anexos_desc = [
    ("Anexo A: Catálogo de Servicios Detallado (Excel)", "anexo_A_catalogo_servicios.xlsx", "Matriz completa con volúmenes transaccionales, canales, tiempos de respuesta, documentos solicitados y justificación técnica para los 10 servicios auditados."),
    ("Anexo B: Diagnóstico de los Ocho Ejes y Líneas Base (Excel)", "anexo_B_diagnostico_ejes.xlsx", "Evaluación de los 23 indicadores de gobierno digital con sus fuentes fechadas, evidencias, niveles y observaciones de campo."),
    ("Anexo C: Gráfico Polar de Madurez de Gobierno Digital (PNG)", "anexo_C_madurez_gobierno_digital.png", "Gráfico de radar en alta definición generado por GD04_madurez.py contrastando el nivel basal 1.17 frente a la meta PETI 3.50."),
    ("Anexo D: Tablero de Metabase y Visualizaciones (PDF)", "anexo_D_tablero_metabase.pdf", "Compilación de 7 páginas en PDF de alta resolución con el master dashboard y las 6 visualizaciones individuales."),
    ("Anexo E: Recorrido del Usuario en Servicios Clave (PDF)", "anexo_E_recorrido_usuario.pdf", "Mapeo detallado de la experiencia de usuario (User Journey Map) en los servicios troncales S-01 Matrícula y S-02 Registro de Actas/Notas."),
    ("Anexo F: Sección 6.1 Oficial del PETI (PDF)", "anexo_F_seccion_6_1.pdf", "Documento normativo de la Sección 6.1 (Situación Actual del Gobierno Digital) listo para su integración en el documento maestro del plan.")
]
for tit, arch, desc in anexos_desc:
    add_styled_paragraph(doc, f"{tit} [Archivo: anexos/{arch}]:\n{desc}", bold_prefix="• ")

# Guardar documento final
doc.save(OUTPUT_DOCX)
print(f"\n[ÉXITO] Documento Word final generado exitosamente: {OUTPUT_DOCX}")
